"""Generate the current system test report from its maintained Markdown source."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from generate_requirements_docx import (
    ROOT,
    add_field,
    configure_page,
    configure_styles,
    render_markdown,
    set_run_font,
)

SOURCE = ROOT / "docs" / "test-report.md"
OUTPUT = ROOT / "docs" / "deliverables" / "下午班-08组-WEMOVE-原网站重构-测试报告-v0.9.docx"


def build() -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc.sections[0])
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Normal"].font.size = Pt(10)

    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("WEMOVE 原网站重构  ·  系统测试报告  ·  v0.9"), size=8, bold=True)

    footer = doc.sections[0].footer
    table = footer.add_table(rows=1, cols=2, width=Cm(16.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left = table.cell(0, 0).paragraphs[0]
    set_run_font(left.add_run("下午班第 8 组  ·  2026-09-08"), size=8, color="5D6975")
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(right.add_run("第 "), size=8, color="5D6975")
    add_field(right, " PAGE ")
    set_run_font(right.add_run(" 页"), size=8, color="5D6975")

    source_text = SOURCE.read_text(encoding="utf-8")
    current_report = source_text.split("## 历史归档", 1)[0].rstrip()
    render_markdown(doc, current_report)
    doc.core_properties.title = "WEMOVE 原网站重构系统测试报告 v0.9"
    doc.core_properties.subject = "当前工作区复核与历史整合测试证据"
    doc.core_properties.author = "甘文韬、龙祖怡"
    doc.core_properties.keywords = "WEMOVE, 原网站重构, 测试报告, 第8组"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
