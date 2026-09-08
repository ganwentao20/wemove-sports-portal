"""Generate the five-member assignment report from its maintained Markdown source."""
from docx import Document
from docx.shared import Inches, Pt
from generate_requirements_docx import ROOT, configure_styles, render_markdown, set_run_font

output = ROOT / 'docs/deliverables/下午班-08组-WEMOVE-原网站重构-分工与工作量报告-v1.3.docx'
doc = Document()
configure_styles(doc)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
doc.styles['Heading 1'].paragraph_format.page_break_before = False
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('第 8 组分工与工作量报告', 0)
render_markdown(doc, (ROOT/'docs/plans/workload-current.md').read_text(encoding='utf-8'))
for table in doc.tables:
    count = len(table.columns)
    widths = [1.1, 2.9, 2.0, .9] if count == 4 else [1.4, 2.7, 2.8]
    if len(widths) == count:
        table.autofit = False
        for column, width in zip(table.columns, widths): column.width = Inches(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs: set_run_font(run, size=9)
doc.core_properties.title = '下午班第8组 WEMOVE 原网站重构分工与工作量报告 v1.3'
doc.core_properties.author = '甘文韬'
doc.core_properties.subject = '五人团队；甘文韬兼任 M1/MB；计划工作量与实际贡献分别统计'
doc.save(output)
print(output)
