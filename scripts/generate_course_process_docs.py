"""Generate the course progress plan and formal minutes from maintained Markdown."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from generate_requirements_docx import (
    ROOT,
    add_field,
    configure_page,
    configure_styles,
    render_markdown,
    set_run_font,
)


OUTPUTS = [
    (
        ROOT / "docs/plans/schedule-current.md",
        ROOT / "docs/deliverables/下午班-08组-WEMOVE-原网站重构-项目进度计划-v1.0.docx",
        "项目进度计划",
        "v1.0",
    ),
    (
        ROOT / "docs/meetings/2026-09-04-meeting-01.md",
        ROOT / "docs/deliverables/下午班-08组-WEMOVE-原网站重构-正式会议纪要-第1次-v1.0.docx",
        "正式会议纪要（第 1 次）",
        "v1.0",
    ),
]


def build(source: Path, output: Path, label: str, version: str) -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc.sections[0])
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.sections[0].top_margin = Cm(1.7)
    doc.sections[0].bottom_margin = Cm(1.6)

    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(
        header.add_run(f"WEMOVE 原网站重构  ·  {label}  ·  {version}"),
        size=8,
        bold=True,
    )

    footer = doc.sections[0].footer
    table = footer.add_table(rows=1, cols=2, width=Cm(16.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left = table.cell(0, 0).paragraphs[0]
    set_run_font(
        left.add_run("软件开发实践2  ·  下午班第 8 组  ·  2026-09-08"),
        size=8,
        color="5D6975",
    )
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(right.add_run("第 "), size=8, color="5D6975")
    add_field(right, " PAGE ")
    set_run_font(right.add_run(" 页"), size=8, color="5D6975")

    render_markdown(doc, source.read_text(encoding="utf-8"))
    doc.core_properties.title = f"WEMOVE 原网站重构{label} {version}"
    doc.core_properties.subject = "课程项目过程材料"
    doc.core_properties.author = "下午班第 8 组"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    for item in OUTPUTS:
        build(*item)
