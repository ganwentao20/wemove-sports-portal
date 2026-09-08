"""Generate the corrected AI summary and five individual course reports."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from generate_requirements_docx import ROOT, configure_styles, render_markdown, set_run_font


OUTPUTS = [
    (
        ROOT / "docs/ai-assistance-summary.md",
        ROOT / "docs/deliverables/下午班-08组-WEMOVE-原网站重构-大模型辅助实践说明-v1.1.docx",
        "WEMOVE 原网站重构五人大模型辅助实践说明 v1.1",
        "第 8 组",
    ),
    *[
        (
            ROOT / f"docs/individual-reports/{name}.md",
            ROOT / f"docs/deliverables/Web开发技术现状报告_{name}_原网站重构校订版.docx",
            f"Web 开发技术现状与课程思政报告（{name}）",
            name,
        )
        for name in ("陈婧琳", "甘文韬", "龙祖怡", "倪依玲", "周慧莹")
    ],
]


def add_footer(doc: Document) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("WEMOVE 原网站重构 · 软件开发实践2 · 2026-09-08 校订")
    set_run_font(run, size=8, color="5D6975")


def generate(source: Path, output: Path, title: str, author: str) -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Normal"].font.size = Pt(10.5)
    render_markdown(doc, source.read_text(encoding="utf-8"))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=8.5)
    add_footer(doc)
    doc.core_properties.title = title
    doc.core_properties.author = author
    doc.core_properties.subject = "保留原网站前端与原图，复用既有后端并补齐课程选题功能"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    for item in OUTPUTS:
        generate(*item)
